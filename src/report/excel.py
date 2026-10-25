"""Excel-отчёт с подсветкой: красный — аномалия, цветовая шкала — доходимость и расхождения."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

INK = "1F2937"
ACCENT = "0F766E"
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)
TOTAL_FILL = PatternFill("solid", fgColor="E6F4F1")
RED_FILL = PatternFill("solid", fgColor="FDE2E1")
RED_FONT = Font(color="B42318", bold=True)
THIN = Side(style="thin", color="D0D5DD")

HEADERS = {
    "date": "Дата",
    "ssp": "SSP",
    "placement": "Площадка",
    "dsp_clicks": "Клики DSP",
    "dsp_clicks_matched": "Клики DSP (дни с Верификатором)",
    "verifier_total": "Кликов всего с учётом IVT",
    "givt": "GIVT",
    "sivt": "SIVT",
    "clean_clicks": "Чистые клики",
    "visits": "Визиты",
    "visits_matched": "Визиты (дни с Верификатором)",
    "robots": "Визиты роботов",
    "add_to_cart": "Цель: корзина",
    "form_submit": "Цель: форма",
    "reach_dsp": "Доходимость от DSP",
    "reach_verifier": "Доходимость от Верификатора",
    "dsp_vs_verifier_pct": "Расхождение DSP / Вериф., %",
    "givt_share": "Доля GIVT",
    "sivt_share": "Доля SIVT",
    "robot_share": "Роботность",
    "cr_cart": "CR корзина",
    "cr_form": "CR форма",
    "avg_daily_reach": "Средняя доходимость по дням",
    "avg_daily_gap_pct": "Среднее расхождение по дням, %",
    "z": "robust z",
    "trend": "Тренд (STL)",
    "flag": "Аномалия",
    "kind": "Тип",
    "slice": "Срез",
    "start": "С",
    "end": "По",
    "days": "Дней",
    "goal": "Цель",
    "reaches": "Достижений",
    "cr": "CR",
    "ci_low": "CI нижн.",
    "ci_high": "CI верхн.",
    "source": "Источник",
    "issue": "Что поправлено",
    "rows": "Строк",
    "detail": "Детали",
}
PERCENT = {
    "reach_dsp",
    "reach_verifier",
    "givt_share",
    "sivt_share",
    "robot_share",
    "cr_cart",
    "cr_form",
    "avg_daily_reach",
    "cr",
    "ci_low",
    "ci_high",
}


def _write_table(ws: Worksheet, df: pd.DataFrame, start_row: int = 1, title: str | None = None) -> tuple[int, int]:
    """Пишет таблицу с шапкой; возвращает (первая строка данных, последняя строка)."""
    row = start_row
    if title:
        ws.cell(row=row, column=1, value=title).font = Font(bold=True, size=13, color=INK)
        row += 1
    cols = list(df.columns)
    for j, col in enumerate(cols, 1):
        c = ws.cell(row=row, column=j, value=HEADERS.get(col, col))
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=THIN)
    ws.row_dimensions[row].height = 32
    first = row + 1
    for i, rec in enumerate(df.itertuples(index=False), first):
        for j, (col, value) in enumerate(zip(cols, rec), 1):
            if isinstance(value, pd.Timestamp):
                value = value.to_pydatetime().date()
            if pd.isna(value) if not isinstance(value, (list, dict)) else False:
                value = None
            c = ws.cell(row=i, column=j, value=value)
            c.border = Border(bottom=THIN)
            if col in PERCENT:
                c.number_format = "0.0%"
            elif col == "date" or col in ("start", "end"):
                c.number_format = "DD.MM.YYYY"
            elif isinstance(value, (int, float)) and col not in (
                "z",
                "trend",
                "dsp_vs_verifier_pct",
                "avg_daily_gap_pct",
            ):
                c.number_format = "#,##0"
    last = first + len(df) - 1
    for j, col in enumerate(cols, 1):
        width = max(len(str(HEADERS.get(col, col))) * 0.9, 11)
        ws.column_dimensions[get_column_letter(j)].width = min(width, 34)
    return first, last


def _col(df: pd.DataFrame, name: str) -> str:
    return get_column_letter(list(df.columns).index(name) + 1)


def _flag_rows(ws: Worksheet, df: pd.DataFrame, first: int, last: int) -> None:
    """Красим строку, где flag = TRUE."""
    flag = _col(df, "flag")
    rng = f"A{first}:{get_column_letter(len(df.columns))}{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"${flag}{first}=TRUE"], fill=RED_FILL, font=RED_FONT))


def _scale(ws: Worksheet, df: pd.DataFrame, col: str, first: int, last: int, reverse: bool = False) -> None:
    low, high = ("F8696B", "63BE7B") if not reverse else ("63BE7B", "F8696B")
    letter = _col(df, col)
    ws.conditional_formatting.add(
        f"{letter}{first}:{letter}{last}",
        ColorScaleRule(
            start_type="min",
            start_color=low,
            mid_type="percentile",
            mid_value=50,
            mid_color="FFEB84",
            end_type="max",
            end_color=high,
        ),
    )


def write(res, path: Path) -> Path:
    wb = Workbook()

    # 1. Сводка
    ws = wb.active
    ws.title = "Сводка"
    ws["A1"] = f"Traffic Quality Lab — кампания {res.cfg['campaign']}"
    ws["A1"].font = Font(bold=True, size=16, color=ACCENT)
    ws["A2"] = f"Период {res.cfg['start']} – {res.cfg['end']}. Данные синтетические."
    ws["A2"].font = Font(italic=True, color="667085")
    cols = [
        "ssp",
        "dsp_clicks",
        "verifier_total",
        "visits",
        "reach_dsp",
        "reach_verifier",
        "avg_daily_reach",
        "dsp_vs_verifier_pct",
        "avg_daily_gap_pct",
        "givt_share",
        "robot_share",
        "cr_cart",
        "cr_form",
    ]
    summary = res.summary[cols]
    first, last = _write_table(ws, summary, start_row=4, title="Итого за период")
    for j in range(1, len(cols) + 1):
        ws.cell(row=last, column=j).fill = TOTAL_FILL
        ws.cell(row=last, column=j).font = Font(bold=True)

    row = last + 3
    windows = res.windows.copy()
    windows["kind"] = windows["kind"].map({"reach": "Провал доходимости", "givt": "Рост GIVT"})
    _, last = _write_table(ws, windows, start_row=row, title="Найденные периоды аномалий")
    row = last + 3
    if not res.goals_over.empty:
        _write_table(ws, res.goals_over, start_row=row, title="Целей больше, чем визитов")
    ws.freeze_panes = "B6"

    # 2. По дням × SSP
    ws = wb.create_sheet("По дням SSP")
    df = res.daily_ssp.merge(res.reach_flags[["date", "ssp", "z", "trend", "flag"]], on=["date", "ssp"])
    df = df[
        [
            "date",
            "ssp",
            "dsp_clicks",
            "verifier_total",
            "visits",
            "reach_dsp",
            "reach_verifier",
            "dsp_vs_verifier_pct",
            "givt_share",
            "robot_share",
            "z",
            "trend",
            "flag",
        ]
    ]
    first, last = _write_table(ws, df)
    _flag_rows(ws, df, first, last)
    _scale(ws, df, "reach_verifier", first, last)
    _scale(ws, df, "dsp_vs_verifier_pct", first, last, reverse=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{last}"

    # 3. Площадки и GIVT
    ws = wb.create_sheet("Площадки GIVT")
    df = res.daily_placement.merge(res.givt_flags[["date", "placement", "z", "flag"]], on=["date", "placement"])
    df = df[["date", "placement", "verifier_total", "givt", "givt_share", "visits", "reach_verifier", "z", "flag"]]
    first, last = _write_table(ws, df)
    _flag_rows(ws, df, first, last)
    _scale(ws, df, "givt_share", first, last, reverse=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{last}"

    # 4. Конверсии
    ws = wb.create_sheet("Конверсии")
    row = 1
    for goal, table in res.cr.items():
        _, last = _write_table(ws, table, start_row=row, title=f"Цель {goal}")
        chi = next(c for c in res.chi2 if c["goal"] == goal)
        ws.cell(row=last + 1, column=1, value=f"Хи-квадрат = {chi['chi2']}, df = {chi['dof']}, p = {chi['p_value']}")
        row = last + 4

    # 5. Витрина целиком
    ws = wb.create_sheet("Витрина")
    _, last = _write_table(ws, res.mart.drop(columns=["users", "bounces"], errors="ignore"))
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{last}"

    # 6. Чистка
    ws = wb.create_sheet("Чистка")
    _write_table(ws, res.cleaning.to_frame())
    ws.column_dimensions["D"].width = 60

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
